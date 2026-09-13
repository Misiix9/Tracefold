from __future__ import annotations

import csv
import io
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TEST_HEADERS = [
    "Record Type", "Run", "Account", "Auth Mode", "Test Case", "HTTP Method", "Exact URL",
    "Endpoint Path", "Expected Status", "Actual Status", "Result", "Duration (ms)",
    "Request Query", "Request Payload", "Response Type", "Response Size (bytes)",
    "Response Body", "Assertions", "Called From", "Found Through", "Notes", "Error", "Timestamp",
]


def json_cell(value: Any) -> str:
    if value is None or value == "":
        return ""
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)


def test_export_rows(result: dict[str, Any]) -> list[list[Any]]:
    run_name = result.get("run", {}).get("name", "")
    rows: list[list[Any]] = []
    for item in result.get("test_results", []):
        assertions = item.get("assertions", [])
        assertion_text = "; ".join(
            f"{entry.get('type')} {entry.get('path', '')}: {'PASS' if entry.get('passed') else 'FAIL'}"
            for entry in assertions
        )
        rows.append([
            "TEST", run_name, item.get("account_name"), item.get("auth_mode"), item.get("test_case_name"),
            item.get("method"), item.get("url"), item.get("endpoint_path"),
            ", ".join(map(str, item.get("expected_statuses", []))), item.get("actual_status"), item.get("result"),
            item.get("duration_ms"), json_cell(item.get("request_query")), json_cell(item.get("request_payload")),
            item.get("response_content_type"), item.get("response_size_bytes"), json_cell(item.get("response_body")),
            assertion_text, "", "test case", item.get("notes"), item.get("error"), item.get("timestamp"),
        ])
    for crawl in result.get("crawl_runs", []):
        account_name = crawl.get("account_name", "")
        auth_mode = crawl.get("auth_mode", "")
        scan = crawl.get("result", {})
        for endpoint in scan.get("endpoints", []):
            locations = []
            for location in endpoint.get("called_from", []):
                value = location.get("page") or location.get("source") or ""
                if location.get("source") and location["source"] != value:
                    value += " | " + location["source"]
                if location.get("line"):
                    value += f":{location['line']}"
                locations.append(value)
            exact_urls = endpoint.get("concrete_urls", [])
            rows.append([
                "DISCOVERY", run_name, account_name, auth_mode, "Full site crawl", endpoint.get("method"),
                "\n".join(exact_urls), endpoint.get("path"), "", ", ".join(map(str, endpoint.get("status_codes", []))),
                "OBSERVED", "", json_cell(endpoint.get("query_parameters")),
                json_cell(endpoint.get("request", {}).get("observed_samples")),
                ", ".join(endpoint.get("request", {}).get("content_types", [])), "",
                json_cell(endpoint.get("responses", [])), "", "\n".join(dict.fromkeys(locations)),
                ", ".join(endpoint.get("discovered_by", [])), "", "", scan.get("scan", {}).get("completed_at"),
            ])
    return rows


def make_test_csv(result: dict[str, Any]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    writer.writerow(TEST_HEADERS)
    writer.writerows(test_export_rows(result))
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def make_test_xlsx(result: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.sheet_view.showGridLines = False
    run = result.get("run", {})
    summary["B2"] = run.get("name", "Endpoint test run")
    summary["B2"].font = Font(name="Arial", size=15, bold=True, color="17191C")
    summary["B3"] = f"Completed {run.get('completed_at', '')}"
    summary["B3"].font = Font(name="Arial", size=10, italic=True, color="6B7178")
    summary["B5"] = "TEST RESULTS"
    summary["B5"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    summary["B5"].fill = PatternFill("solid", fgColor="17191C")
    summary.merge_cells("B5:E5")
    metrics = [
        ("Total", run.get("tests", 0), "E7EEF6"),
        ("Passed", run.get("passed", 0), "E7F5E5"),
        ("Failed", run.get("failed", 0), "FCE8E8"),
        ("Errors", run.get("errors", 0), "FFF1D7"),
    ]
    for offset, (label, value, color) in enumerate(metrics, start=2):
        summary.cell(7, offset, label)
        summary.cell(8, offset, value)
        summary.cell(7, offset).font = Font(name="Arial", size=10, color="666C73")
        summary.cell(8, offset).font = Font(name="Arial", size=14, bold=True, color="17191C")
        summary.cell(8, offset).fill = PatternFill("solid", fgColor=color)
        summary.cell(8, offset).alignment = Alignment(horizontal="center")
    summary["B11"] = "ACCOUNTS"
    summary["B11"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    summary["B11"].fill = PatternFill("solid", fgColor="17191C")
    summary.merge_cells("B11:E11")
    summary.append([])
    row = 13
    summary.cell(row, 2, "Name")
    summary.cell(row, 3, "Auth mode")
    _style_header(summary, row, 2, 3)
    for account in run.get("accounts", []):
        row += 1
        summary.cell(row, 2, account.get("name"))
        summary.cell(row, 3, account.get("auth_mode"))
    summary.column_dimensions["A"].width = 3
    summary.column_dimensions["B"].width = 28
    summary.column_dimensions["C"].width = 22
    summary.column_dimensions["D"].width = 18
    summary.column_dimensions["E"].width = 18

    results_sheet = workbook.create_sheet("Test Results")
    _write_data_sheet(results_sheet, TEST_HEADERS, test_export_rows(result), freeze="F5")
    result_column = TEST_HEADERS.index("Result") + 1
    first_data_row = 5
    last_data_row = max(first_data_row, results_sheet.max_row)
    full_range = f"A{first_data_row}:W{last_data_row}"
    result_letter = get_column_letter(result_column)
    results_sheet.conditional_formatting.add(full_range, FormulaRule(
        formula=[f'=${result_letter}{first_data_row}="PASS"'], fill=PatternFill("solid", fgColor="E8F5E6")
    ))
    results_sheet.conditional_formatting.add(full_range, FormulaRule(
        formula=[f'=${result_letter}{first_data_row}="FAIL"'], fill=PatternFill("solid", fgColor="FCE8E8")
    ))
    results_sheet.conditional_formatting.add(full_range, FormulaRule(
        formula=[f'=${result_letter}{first_data_row}="ERROR"'], fill=PatternFill("solid", fgColor="FFF1D7")
    ))

    api_headers = [
        "Account", "Method", "Exact URL", "Normalized Path", "Status Codes", "Payload Seen",
        "Payload Schema", "Called From", "Found Through", "Authentication",
    ]
    api_rows: list[list[Any]] = []
    page_rows: list[list[Any]] = []
    issue_rows: list[list[Any]] = []
    for crawl in result.get("crawl_runs", []):
        account = crawl.get("account_name", "")
        scan = crawl.get("result", {})
        for endpoint in scan.get("endpoints", []):
            locations = [item.get("page") or item.get("source") or "" for item in endpoint.get("called_from", [])]
            api_rows.append([
                account, endpoint.get("method"), "\n".join(endpoint.get("concrete_urls", [])), endpoint.get("path"),
                ", ".join(map(str, endpoint.get("status_codes", []))),
                "Yes" if endpoint.get("request", {}).get("has_body") else "No",
                json_cell(endpoint.get("request", {}).get("schema")), "\n".join(dict.fromkeys(locations)),
                ", ".join(endpoint.get("discovered_by", [])), endpoint.get("authentication"),
            ])
        for page in scan.get("pages", []):
            page_rows.append([account, page.get("status"), page.get("url"), page.get("title"), page.get("depth"), page.get("error")])
        issue_rows.extend([[account, "Warning", item] for item in scan.get("warnings", [])])
        issue_rows.extend([[account, "Error", item] for item in scan.get("errors", [])])
    api_sheet = workbook.create_sheet("Discovered APIs")
    _write_data_sheet(api_sheet, api_headers, api_rows, freeze="C5")
    pages_sheet = workbook.create_sheet("Crawled Pages")
    _write_data_sheet(pages_sheet, ["Account", "Status", "Exact Page URL", "Title", "Depth", "Error"], page_rows, freeze="C5")
    issues_sheet = workbook.create_sheet("Issues")
    _write_data_sheet(issues_sheet, ["Account", "Type", "Message"], issue_rows, freeze="A5")

    widths = {
        "Test Results": [14, 22, 20, 14, 28, 13, 58, 40, 18, 15, 13, 15, 28, 38, 22, 20, 65, 40, 60, 22, 35, 45, 24],
        "Discovered APIs": [22, 12, 60, 42, 18, 16, 65, 65, 24, 24],
        "Crawled Pages": [22, 12, 70, 40, 10, 60],
        "Issues": [22, 14, 100],
    }
    for sheet_name, values in widths.items():
        sheet = workbook[sheet_name]
        for index, width in enumerate(values, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_data_sheet(sheet: Any, headers: list[str], rows: list[list[Any]], freeze: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A2"] = sheet.title
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color="17191C")
    header_row = 4
    for column, value in enumerate(headers, start=1):
        sheet.cell(header_row, column, value)
    _style_header(sheet, header_row, 1, len(headers))
    for row_values in rows:
        sheet.append(row_values)
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, sheet.max_row)}"
    thin = Side(style="thin", color="D8DCE0")
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="25282C")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)


def _style_header(sheet: Any, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = sheet.cell(row, column)
        cell.fill = PatternFill("solid", fgColor="17191C")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 28
