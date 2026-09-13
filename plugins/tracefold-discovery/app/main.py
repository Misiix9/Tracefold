from __future__ import annotations

import asyncio
import csv
import io
import json
import logging
import os
import re
import subprocess
import sys
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from starlette.middleware.trustedhost import TrustedHostMiddleware
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .auth import SessionCaptureManager, account_auth_context
from .exporters import make_test_csv, make_test_xlsx
from .models import AccountInput, ScanConfig, ScanJob, TestCaseInput, TestRunInput
from .scanner import SiteScanner
from .storage import AccountStore, TestCaseStore
from .test_runner import TestRunner


ROOT = Path(__file__).resolve().parent.parent
STATIC = ROOT / "static"


def resolve_data_dir() -> Path:
    """Settings, accounts, sessions, cases and results live outside the installed code.

    Tracefold supplies a persistent per-plugin directory, so replacing the plugin package
    on update never touches user data. Standalone runs keep the previous layout.
    """
    supplied = os.environ.get("TRACEFOLD_DISCOVERY_DATA_DIR") or os.environ.get("TRACEFOLD_PLUGIN_DATA_DIR")
    if supplied:
        directory = Path(supplied).expanduser().resolve()
        directory.mkdir(parents=True, exist_ok=True)
        return directory
    return ROOT / ".tracefold"


DATA = resolve_data_dir()
LEGACY_DATA = ROOT / "data"
STANDALONE_DATA = ROOT / ".tracefold"
# Carry forward data from an earlier standalone install exactly once, without ever
# overwriting something that already exists in the destination.
for legacy in (STANDALONE_DATA, LEGACY_DATA):
    if legacy.resolve() == DATA.resolve() or not legacy.exists():
        continue
    DATA.mkdir(parents=True, exist_ok=True)
    for child in legacy.iterdir():
        destination = DATA / child.name
        if destination.exists():
            continue
        if child.is_dir():
            import shutil

            shutil.copytree(child, destination)
        elif child.is_file():
            destination.write_bytes(child.read_bytes())
SCAN_DATA = DATA / "scans"
TEST_RUN_DATA = DATA / "test-runs"
SETTINGS_FILE = DATA / "settings.json"
SCAN_DATA.mkdir(parents=True, exist_ok=True)
TEST_RUN_DATA.mkdir(parents=True, exist_ok=True)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logger = logging.getLogger("tracefold")

APP_VERSION = os.environ.get("TRACEFOLD_PLUGIN_VERSION") or "2.1.0"
app = FastAPI(title="Tracefold Discovery", version=APP_VERSION, docs_url=None, redoc_url=None)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=["127.0.0.1", "localhost"])
jobs: dict[str, ScanJob] = {}
test_jobs: dict[str, ScanJob] = {}
tasks: set[asyncio.Task[Any]] = set()
account_store = AccountStore(DATA)
test_case_store = TestCaseStore(DATA)
session_captures = SessionCaptureManager(account_store)


@app.middleware("http")
async def security_headers(request: Request, call_next: Any) -> Response:
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Cache-Control"] = "no-store"
    # X-Frame-Options cannot express an allowed embedder, so framing is governed by
    # frame-ancestors alone: the Tracefold desktop window, and nothing else.
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; style-src 'self'; script-src 'self'; img-src 'self' data:; "
        "connect-src 'self'; base-uri 'none'; form-action 'self'; "
        "frame-ancestors 'self' tauri://localhost http://tauri.localhost https://tauri.localhost"
    )
    return response


@app.get("/")
async def index() -> FileResponse:
    return FileResponse(STATIC / "index.html")


@app.get("/api/health")
async def health() -> dict[str, str]:
    return {"status": "ok", "version": APP_VERSION, "vault": account_store.vault.backend}


DEFAULT_SETTINGS = {
    "max_pages": 50, "max_depth": 5, "request_delay_ms": 150, "navigation_timeout_ms": 20000,
    "include_subdomains": False, "ignore_https_errors": False, "discover_openapi": True, "safe_probing": True,
    "normalize_dynamic_ids": True, "max_route_variants": 1, "custom_id_patterns": [],
    "ignore_path_patterns": [], "normalize_query_values": True,
}

def load_settings() -> dict[str, Any]:
    if not SETTINGS_FILE.exists():
        return dict(DEFAULT_SETTINGS)
    try:
        value = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
        return {**DEFAULT_SETTINGS, **value}
    except (OSError, json.JSONDecodeError):
        return dict(DEFAULT_SETTINGS)

def save_settings(value: dict[str, Any]) -> dict[str, Any]:
    SETTINGS_FILE.parent.mkdir(parents=True, exist_ok=True)
    merged = {**DEFAULT_SETTINGS, **value}
    temporary = SETTINGS_FILE.with_suffix(".tmp")
    temporary.write_text(json.dumps(merged, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(SETTINGS_FILE)
    return merged

@app.get("/api/settings")
async def get_settings() -> dict[str, Any]:
    return load_settings()

@app.put("/api/settings")
async def put_settings(value: dict[str, Any]) -> dict[str, Any]:
    try:
        validated = ScanConfig(target_url="https://example.invalid", **{k: v for k, v in value.items() if k in DEFAULT_SETTINGS})
    except Exception as exc:
        raise HTTPException(422, safe_error(exc))
    safe = {key: getattr(validated, key) for key in DEFAULT_SETTINGS}
    return save_settings(safe)

@app.get("/api/accounts")
async def list_accounts() -> list[dict[str, Any]]:
    return account_store.list_public()


@app.post("/api/accounts", status_code=201)
async def create_account(value: AccountInput) -> dict[str, Any]:
    try:
        return account_store.create(value)
    except ValueError as exc:
        raise HTTPException(422, str(exc))


@app.put("/api/accounts/{account_id}")
async def update_account(account_id: str, value: AccountInput) -> dict[str, Any]:
    try:
        return account_store.update(account_id, value)
    except KeyError:
        raise HTTPException(404, "Account not found")
    except ValueError as exc:
        raise HTTPException(422, str(exc))


@app.delete("/api/accounts/{account_id}", status_code=204)
async def delete_account(account_id: str) -> Response:
    await session_captures.cancel(account_id)
    try:
        account_store.delete(account_id)
    except KeyError:
        raise HTTPException(404, "Account not found")
    return Response(status_code=204)


@app.post("/api/accounts/{account_id}/session/start", status_code=202)
async def start_account_session(account_id: str) -> dict[str, str]:
    try:
        return await session_captures.start(account_id)
    except KeyError:
        raise HTTPException(404, "Account not found")
    except Exception as exc:
        raise HTTPException(500, safe_error(exc))


@app.post("/api/accounts/{account_id}/session/finish")
async def finish_account_session(account_id: str) -> dict[str, Any]:
    try:
        return await session_captures.finish(account_id)
    except KeyError:
        raise HTTPException(409, "No login window is open for this account")
    except Exception as exc:
        raise HTTPException(500, safe_error(exc))


@app.post("/api/accounts/{account_id}/session/cancel", status_code=204)
async def cancel_account_session(account_id: str) -> Response:
    await session_captures.cancel(account_id)
    return Response(status_code=204)


@app.get("/api/test-cases")
async def list_test_cases() -> list[dict[str, Any]]:
    return test_case_store.list()


@app.post("/api/test-cases", status_code=201)
async def create_test_case(value: TestCaseInput) -> dict[str, Any]:
    return test_case_store.create(value)


@app.put("/api/test-cases/{case_id}")
async def update_test_case(case_id: str, value: TestCaseInput) -> dict[str, Any]:
    try:
        return test_case_store.update(case_id, value)
    except KeyError:
        raise HTTPException(404, "Test case not found")


@app.delete("/api/test-cases/{case_id}", status_code=204)
async def delete_test_case(case_id: str) -> Response:
    try:
        test_case_store.delete(case_id)
    except KeyError:
        raise HTTPException(404, "Test case not found")
    return Response(status_code=204)


@app.post("/api/test-runs", status_code=202)
async def create_test_run(value: TestRunInput) -> dict[str, str]:
    job_id = uuid.uuid4().hex
    job = ScanJob(id=job_id)
    test_jobs[job_id] = job
    task = asyncio.create_task(run_test_job(job_id, value))
    tasks.add(task)
    task.add_done_callback(tasks.discard)
    return {"id": job_id, "status": "queued"}


@app.get("/api/test-runs/{job_id}")
async def get_test_run(job_id: str) -> dict[str, Any]:
    job = get_test_job(job_id)
    return job.model_dump(exclude={"cancel_requested"})


@app.post("/api/test-runs/{job_id}/cancel", status_code=202)
async def cancel_test_run(job_id: str) -> dict[str, str]:
    job = get_test_job(job_id)
    if job.status not in {"queued", "running"}:
        raise HTTPException(409, "This test run is no longer running")
    job.cancel_requested = True
    job.phase = "Cancelling"
    return {"status": "cancelling"}


@app.get("/api/test-runs/{job_id}/export/{format_name}")
async def export_test_run(job_id: str, format_name: str) -> Response:
    job = get_test_job(job_id)
    if job.status != "completed" or not job.result:
        raise HTTPException(409, "The test run has not completed")
    if format_name == "json":
        content = json.dumps(job.result, ensure_ascii=False, indent=2).encode("utf-8")
        return Response(content, media_type="application/json", headers={"Content-Disposition": 'attachment; filename="tracefold-test-run.json"'})
    if format_name == "csv":
        return Response(make_test_csv(job.result), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="tracefold-test-run.csv"'})
    if format_name == "xlsx":
        return StreamingResponse(
            io.BytesIO(make_test_xlsx(job.result)),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="tracefold-test-run.xlsx"'},
        )
    raise HTTPException(404, "Supported formats: json, csv, xlsx")


EXPORTS = DATA / "exports"
EXPORT_MEDIA = {
    "json": "application/json",
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def render_export(result: dict[str, Any], format_name: str, kind: str) -> bytes:
    if format_name == "json":
        return json.dumps(result, ensure_ascii=False, indent=2).encode("utf-8")
    if format_name == "csv":
        return make_csv(result) if kind == "scan" else make_test_csv(result)
    if format_name == "xlsx":
        return make_xlsx(result) if kind == "scan" else make_test_xlsx(result)
    raise HTTPException(404, "Supported formats: json, csv, xlsx")


def write_export(result: dict[str, Any], format_name: str, kind: str, stem: str) -> dict[str, str]:
    """Save a report next to the plugin's other data.

    A hosted plugin is displayed inside Tracefold's window, where a browser-style download
    has nowhere to go. Writing the file and reporting its location works identically in the
    host, in a standalone browser, and with no user interaction at all.
    """
    if format_name not in EXPORT_MEDIA:
        raise HTTPException(404, "Supported formats: json, csv, xlsx")
    content = render_export(result, format_name, kind)
    EXPORTS.mkdir(parents=True, exist_ok=True)
    safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-.") or kind
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    destination = EXPORTS / f"{safe_stem}-{stamp}.{format_name}"
    temporary = destination.with_name(destination.name + ".part")
    temporary.write_bytes(content)
    temporary.replace(destination)
    return {"path": str(destination), "name": destination.name, "bytes": str(len(content))}


@app.post("/api/scans/{job_id}/save/{format_name}")
async def save_scan_export(job_id: str, format_name: str) -> dict[str, str]:
    job = get_job(job_id)
    if job.status != "completed" or not job.result:
        raise HTTPException(409, "The scan has not completed")
    host = urlparse(str(job.result.get("scan", {}).get("target_url", ""))).netloc or "scan"
    return await asyncio.to_thread(write_export, job.result, format_name, "scan", f"tracefold-{host}")


@app.post("/api/test-runs/{job_id}/save/{format_name}")
async def save_test_export(job_id: str, format_name: str) -> dict[str, str]:
    job = get_test_job(job_id)
    if job.status != "completed" or not job.result:
        raise HTTPException(409, "The test run has not completed")
    return await asyncio.to_thread(write_export, job.result, format_name, "test-run", "tracefold-test-run")


@app.post("/api/reveal")
async def reveal_export(payload: dict[str, Any]) -> dict[str, str]:
    """Show a saved report in the desktop file manager.

    Only files this plugin wrote into its own exports directory can be revealed, so a
    crafted request cannot use it to open arbitrary locations.
    """
    raw = str(payload.get("path", ""))
    try:
        target = Path(raw).resolve()
        target.relative_to(EXPORTS.resolve())
    except (OSError, ValueError):
        raise HTTPException(422, "Only saved Discovery reports can be revealed")
    if not target.is_file():
        raise HTTPException(404, "That report no longer exists")
    try:
        if sys.platform == "darwin":
            subprocess.Popen(["open", "-R", str(target)])
        elif os.name == "nt":
            subprocess.Popen(["explorer", "/select,", str(target)])
        else:
            subprocess.Popen(["xdg-open", str(target.parent)])
    except OSError as exc:
        raise HTTPException(500, safe_error(exc))
    return {"status": "ok", "path": str(target)}


@app.post("/api/scans", status_code=202)
async def create_scan(config: ScanConfig) -> dict[str, str]:
    save_settings({key: getattr(config, key) for key in DEFAULT_SETTINGS if hasattr(config, key)})
    if config.account_id:
        try:
            account = account_store.get_public(config.account_id)
        except KeyError:
            raise HTTPException(404, "Saved account not found")
        if not account_target_allowed(str(config.target_url), account):
            raise HTTPException(422, "The scan target is outside this account's base and allowed API origins")
    job_id = uuid.uuid4().hex
    job = ScanJob(id=job_id)
    jobs[job_id] = job
    task = asyncio.create_task(run_scan(job_id, config))
    tasks.add(task)
    task.add_done_callback(tasks.discard)
    return {"id": job_id, "status": "queued"}


@app.get("/api/scans")
async def list_scans() -> list[dict[str, Any]]:
    values = sorted(jobs.values(), key=lambda item: item.started_at or "", reverse=True)
    return [job.model_dump(exclude={"result", "cancel_requested"}) for job in values[:25]]


@app.get("/api/scans/{job_id}")
async def get_scan(job_id: str) -> dict[str, Any]:
    job = get_job(job_id)
    return job.model_dump(exclude={"cancel_requested"})


@app.post("/api/scans/{job_id}/cancel", status_code=202)
async def cancel_scan(job_id: str) -> dict[str, str]:
    job = get_job(job_id)
    if job.status not in {"queued", "running"}:
        raise HTTPException(409, "This scan is no longer running")
    job.cancel_requested = True
    job.phase = "Cancelling"
    return {"status": "cancelling"}


@app.get("/api/scans/{job_id}/export/{format_name}")
async def export_scan(job_id: str, format_name: str) -> Response:
    job = get_job(job_id)
    if job.status != "completed" or not job.result:
        raise HTTPException(409, "The scan has not completed")
    safe_host = (job.result["scan"].get("target_url") or "scan").split("//")[-1].split("/")[0].replace(":", "_")
    if format_name == "json":
        content = json.dumps(job.result, ensure_ascii=False, indent=2).encode("utf-8")
        return Response(content, media_type="application/json", headers={"Content-Disposition": f'attachment; filename="tracefold-{safe_host}.json"'})
    if format_name == "csv":
        content = make_csv(job.result)
        return Response(content, media_type="text/csv; charset=utf-8", headers={"Content-Disposition": f'attachment; filename="tracefold-{safe_host}.csv"'})
    if format_name == "xlsx":
        content = make_xlsx(job.result)
        return StreamingResponse(
            io.BytesIO(content),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="tracefold-{safe_host}.xlsx"'},
        )
    raise HTTPException(404, "Supported formats: json, csv, xlsx")


async def run_scan(job_id: str, config: ScanConfig) -> None:
    job = jobs[job_id]
    job.status = "running"
    job.started_at = job.now()

    def progress(phase: str, percent: int, pages: int, endpoints: int) -> None:
        job.phase = phase
        job.progress = max(job.progress, min(99, percent))
        job.pages_visited = pages
        job.endpoints_found = endpoints

    auth_context = None
    if config.account_id:
        account = account_store.get_runtime(config.account_id)
        auth_context = account_auth_context(account)
        config = config.model_copy(update={
            "bearer_token": auth_context.get("bearer_token", ""),
            "extra_headers": {**auth_context.get("extra_headers", {}), **config.extra_headers},
        })
    scanner = SiteScanner(config, on_progress=progress, is_cancelled=lambda: job.cancel_requested, auth_context=auth_context)
    try:
        result = await scanner.run()
        job.result = result
        job.pages_visited = result["scan"]["pages_visited"]
        job.endpoints_found = result["scan"]["endpoints_found"]
        job.status = "cancelled" if job.cancel_requested else "completed"
        job.phase = "Cancelled" if job.cancel_requested else "Complete"
        job.progress = 100
        if job.status == "completed":
            await asyncio.to_thread(save_result, job_id, result)
    except Exception as exc:
        logger.exception("Scan %s failed", job_id)
        job.status = "failed"
        job.phase = "Failed"
        job.error = safe_error(exc)
    finally:
        job.completed_at = job.now()


async def run_test_job(job_id: str, request: TestRunInput) -> None:
    job = test_jobs[job_id]
    job.status = "running"
    job.started_at = job.now()

    def progress(phase: str, percent: int, steps: int, results: int) -> None:
        job.phase = phase
        job.progress = max(job.progress, min(99, percent))
        job.pages_visited = steps
        job.endpoints_found = results

    runner = TestRunner(request, account_store, test_case_store, progress, lambda: job.cancel_requested)
    try:
        result = await runner.run()
        job.result = result
        job.status = "cancelled" if job.cancel_requested else "completed"
        job.phase = "Cancelled" if job.cancel_requested else "Complete"
        job.progress = 100
        job.pages_visited = result["run"]["tests"]
        job.endpoints_found = result["run"]["passed"]
        if job.status == "completed":
            await asyncio.to_thread(save_test_result, job_id, result)
    except Exception as exc:
        logger.exception("Test run %s failed", job_id)
        job.status = "failed"
        job.phase = "Failed"
        job.error = safe_error(exc)
    finally:
        job.completed_at = job.now()


def save_result(job_id: str, result: dict[str, Any]) -> None:
    path = SCAN_DATA / f"{job_id}.json"
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


def save_test_result(job_id: str, result: dict[str, Any]) -> None:
    path = TEST_RUN_DATA / f"{job_id}.json"
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


def get_job(job_id: str) -> ScanJob:
    if not job_id.isalnum() or job_id not in jobs:
        raise HTTPException(404, "Scan not found")
    return jobs[job_id]


def get_test_job(job_id: str) -> ScanJob:
    if not job_id.isalnum() or job_id not in test_jobs:
        raise HTTPException(404, "Test run not found")
    return test_jobs[job_id]


def safe_error(exc: Exception) -> str:
    text = str(exc).replace("\n", " ")
    import re

    return re.sub(r"Bearer\s+[A-Za-z0-9._~+/-]+", "Bearer [redacted]", text, flags=re.I)[:1000]


def account_target_allowed(url: str, account: dict[str, Any]) -> bool:
    target = urlparse(url)
    for allowed in [account.get("base_url", ""), *account.get("allowed_origins", [])]:
        origin = urlparse(str(allowed))
        if target.scheme == origin.scheme and target.netloc.lower() == origin.netloc.lower():
            return True
    return False


def endpoint_rows(result: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    for endpoint in result.get("endpoints", []):
        called_from = []
        for item in endpoint.get("called_from", []):
            location = item.get("page") or item.get("source") or ""
            if item.get("source") and item.get("source") != location:
                location += f" | {item['source']}"
            if item.get("line"):
                location += f":{item['line']}"
            called_from.append(location)
        rows.append([
            endpoint.get("method", ""),
            endpoint.get("origin", ""),
            endpoint.get("path", ""),
            ", ".join(str(code) for code in endpoint.get("status_codes", [])),
            "Yes" if endpoint.get("request", {}).get("has_body") else "No",
            ", ".join(endpoint.get("request", {}).get("content_types", [])),
            json.dumps(endpoint.get("request", {}).get("schema", {}), ensure_ascii=False),
            "\n".join(dict.fromkeys(called_from)),
            ", ".join(endpoint.get("discovered_by", [])),
            endpoint.get("authentication", ""),
        ])
    return rows


EXPORT_HEADERS = [
    "Method", "Origin", "Path", "Status codes", "Has payload", "Request content types",
    "Payload schema", "Called from", "Discovered by", "Authentication",
]


def make_csv(result: dict[str, Any]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(EXPORT_HEADERS)
    writer.writerows(endpoint_rows(result))
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def make_xlsx(result: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Endpoints"
    sheet.append(EXPORT_HEADERS)
    for row in endpoint_rows(result):
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="171717")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    widths = [12, 28, 42, 16, 14, 24, 60, 70, 22, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    pages = workbook.create_sheet("Pages")
    pages.append(["URL", "Status", "Title", "Depth", "Error"])
    for page in result.get("pages", []):
        pages.append([page.get("url"), page.get("status"), page.get("title"), page.get("depth"), page.get("error")])
    for cell in pages[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    pages.freeze_panes = "A2"
    pages.auto_filter.ref = pages.dimensions
    for column, width in zip("ABCDE", [70, 12, 40, 10, 70]):
        pages.column_dimensions[column].width = width

    summary = workbook.create_sheet("Summary", 0)
    scan = result.get("scan", {})
    summary.append(["Tracefold Discovery report"])
    summary["A1"].font = Font(size=18, bold=True)
    for label, value in [
        ("Target", scan.get("target_url")),
        ("Started", scan.get("started_at")),
        ("Completed", scan.get("completed_at")),
        ("Pages visited", scan.get("pages_visited")),
        ("Endpoints found", scan.get("endpoints_found")),
        ("Bearer token used", scan.get("bearer_token_used")),
        ("Active mutation probing", scan.get("active_mutation_probing")),
    ]:
        summary.append([label, value])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 80

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


app.mount("/static", StaticFiles(directory=STATIC), name="static")
