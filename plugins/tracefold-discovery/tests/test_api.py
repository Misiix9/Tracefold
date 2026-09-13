from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import account_target_allowed, app, make_csv, make_xlsx


SAMPLE_RESULT = {
    "scan": {
        "target_url": "https://example.com",
        "started_at": "2026-01-01T00:00:00Z",
        "completed_at": "2026-01-01T00:00:01Z",
        "pages_visited": 1,
        "endpoints_found": 1,
        "bearer_token_used": True,
        "active_mutation_probing": False,
    },
    "pages": [{"url": "https://example.com/tickets", "status": 200, "title": "Tickets", "depth": 0, "error": None}],
    "endpoints": [{
        "method": "GET",
        "origin": "https://example.com",
        "path": "/api/tickets/{id}",
        "status_codes": [200],
        "request": {"has_body": False, "content_types": [], "schema": {}},
        "called_from": [{"page": "https://example.com/tickets"}],
        "discovered_by": ["runtime"],
        "authentication": "Bearer supplied",
    }],
}


def test_health_and_security_headers() -> None:
    client = TestClient(app, base_url="http://127.0.0.1")
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"
    # Framing is governed by frame-ancestors alone: X-Frame-Options cannot name an
    # allowed embedder, and sending DENY would stop Tracefold from displaying the plugin.
    assert "x-frame-options" not in {name.lower() for name in response.headers}
    assert "default-src 'self'" in response.headers["content-security-policy"]
    assert "frame-ancestors 'self' tauri://localhost" in response.headers["content-security-policy"]


def test_csv_export_uses_excel_friendly_utf8() -> None:
    content = make_csv(SAMPLE_RESULT)
    assert content.startswith(b"\xef\xbb\xbf")
    assert b"/api/tickets/{id}" in content


def test_xlsx_export_contains_summary_and_rows() -> None:
    content = make_xlsx(SAMPLE_RESULT)
    workbook = load_workbook(BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["Summary", "Endpoints", "Pages"]
    assert workbook["Endpoints"]["C2"].value == "/api/tickets/{id}"
    assert workbook["Pages"]["A2"].value == "https://example.com/tickets"


def test_saved_account_credentials_are_scoped_to_allowed_origins() -> None:
    account = {"base_url": "https://uat.example.com", "allowed_origins": ["https://api.uat.example.com"]}
    assert account_target_allowed("https://uat.example.com/tickets", account)
    assert account_target_allowed("https://api.uat.example.com/v1/tickets", account)
    assert not account_target_allowed("https://evil.example.com", account)
