from __future__ import annotations

import asyncio
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO

from openpyxl import load_workbook

from app.exporters import make_test_csv, make_test_xlsx
from app.models import AccountInput, TestCaseInput as CaseInput, TestRunInput as RunInput
from app.storage import AccountStore, TestCaseStore as CaseStore
from app.test_runner import TestRunner as Runner


class ApiHandler(BaseHTTPRequestHandler):
    def do_GET(self) -> None:  # noqa: N802
        token = self.headers.get("authorization", "")
        status = 403 if token == "Bearer token-a" else 200
        body = json.dumps({"access": status == 200, "path": self.path}).encode()
        self.send_response(status)
        self.send_header("content-type", "application/json")
        self.send_header("content-length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format: str, *args: object) -> None:
        return


def test_runs_one_case_with_two_named_accounts_and_exports(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("app.storage.keyring.get_password", lambda *args: (_ for _ in ()).throw(RuntimeError()))
    server = ThreadingHTTPServer(("127.0.0.1", 0), ApiHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        base_url = f"http://127.0.0.1:{server.server_port}"
        accounts = AccountStore(tmp_path)
        account_a = accounts.create(AccountInput(name="Restricted A", base_url=base_url, auth_mode="bearer", bearer_token="token-a"))
        account_b = accounts.create(AccountInput(name="Allowed B", base_url=base_url, auth_mode="bearer", bearer_token="token-b"))
        cases = CaseStore(tmp_path)
        case = cases.create(CaseInput(
            name="Direct object access",
            method="GET",
            url="/api/Tickets/${ticketId}",
            expected_statuses=[200],
            account_expectations={account_a["id"]: [403], account_b["id"]: [200]},
        ))
        metadata = accounts.collection.read()
        for item in metadata:
            item["variables"] = {"ticketId": "42"}
        accounts.collection.write(metadata)
        request = RunInput(
            name="Authorization check",
            account_ids=[account_a["id"], account_b["id"]],
            test_case_ids=[case["id"]],
        )
        result = asyncio.run(Runner(request, accounts, cases).run())
    finally:
        server.shutdown()
        server.server_close()

    assert [item["account_name"] for item in result["test_results"]] == ["Restricted A", "Allowed B"]
    assert [item["actual_status"] for item in result["test_results"]] == [403, 200]
    assert all(item["result"] == "PASS" for item in result["test_results"])
    assert "token-a" not in str(result)
    assert "Restricted A" in make_test_csv(result).decode("utf-8-sig")

    workbook = load_workbook(BytesIO(make_test_xlsx(result)), read_only=False)
    assert workbook.sheetnames == ["Summary", "Test Results", "Discovered APIs", "Crawled Pages", "Issues"]
    assert workbook["Test Results"]["C5"].value == "Restricted A"
    rules = [rule for group in workbook["Test Results"].conditional_formatting for rule in group.rules]
    assert len(rules) == 3


def test_jsonpath_and_empty_array_assertions() -> None:
    assertions = [
        {"type": "empty_array", "path": "$.private", "value": None},
        {"type": "jsonpath_equals", "path": "$.visible", "value": True},
        {"type": "not_contains", "path": "", "value": "secret"},
    ]
    body = {"private": [], "visible": True}
    result = Runner._evaluate_assertions(assertions, json.dumps(body), body)
    assert all(item["passed"] for item in result)
